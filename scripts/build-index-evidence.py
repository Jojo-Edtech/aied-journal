"""Join dated public index lists to the existing catalog by current ISSN only.

Run with openpyxl available. Source files are local inputs, never fetched implicitly.
Does not copy third-party metrics or infer index membership from JCR metrics.
"""
import argparse
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import re

import openpyxl

SCOPUS_URL = 'https://downloads.ctfassets.net/o78em1y1w4i4/7xtaTxNiNcWRTeZkV86eNy/69cf2d506c905dc299531fdc93049dbb/ext_list_Aug_2026.xlsx'
WOS_URL = 'https://uefiscdi.gov.ro/resource-832705-AIS.JCR2025.iunie2026.xlsx'


def issn(value):
    value = re.sub('[^0-9X]', '', str(value or '').upper())
    return value if re.fullmatch('[0-9]{7}[0-9X]', value) else ''


def build(journals, scopus_path, wos_path):
    scopus = openpyxl.load_workbook(scopus_path, read_only=True, data_only=True)
    scopus_by_issn = defaultdict(list)
    discontinued_ids = {str(row[0]) for row in scopus.worksheets[2].iter_rows(min_row=3, values_only=True) if row[0]}
    for row_number, row in enumerate(scopus.worksheets[0].iter_rows(min_row=2, values_only=True), 2):
        for identifier in {issn(row[2]), issn(row[3])} - {''}:
            scopus_by_issn[identifier].append((row_number, row))
    wos = openpyxl.load_workbook(wos_path, read_only=True, data_only=True)
    wos_by_issn = defaultdict(list)
    for row_number, row in enumerate(wos.worksheets[0].iter_rows(min_row=2, values_only=True), 2):
        for identifier in {issn(row[1]), issn(row[2])} - {''}:
            wos_by_issn[identifier].append((row_number, row))
    entries = {}
    language_names = {'ENG': 'English', 'CHI': 'Chinese', 'ZHO': 'Chinese', 'FRE': 'French', 'GER': 'German', 'SPA': 'Spanish', 'POR': 'Portuguese', 'RUS': 'Russian', 'TUR': 'Turkish', 'JPN': 'Japanese', 'ITA': 'Italian', 'DUT': 'Dutch'}
    ambiguous = []
    for journal in journals:
        historical = {issn(v) for v in journal.get('historical_issns', [])}
        identifiers = ({issn(v) for v in [journal.get('issn'), journal.get('eissn'), *journal.get('issns', [])]} - {''}) - historical
        scopus_matches = {n: row for identifier in identifiers for n, row in scopus_by_issn.get(identifier, [])}
        active = [(n, r) for n, r in scopus_matches.items() if r[4] == 'Active' and r[12] == 'Journal' and not r[6] and str(r[0]) not in discontinued_ids]
        if len({str(r[0]) for _, r in active}) > 1:
            ambiguous.append(journal['id'])
            continue
        indexes, languages, proofs, language_proofs = set(), set(), [], []
        # Language is explicit metadata even on an inactive source. It does not
        # establish current Scopus coverage; only the active loop below does.
        if len({str(row[0]) for row in scopus_matches.values()}) == 1:
            for row_number, row in scopus_matches.items():
                langs = [language_names.get(v.strip(), v.strip()) for v in str(row[7] or '').split(';') if v.strip()]
                languages.update(langs)
                if langs:
                    language_proofs.append({'source': 'Elsevier Scopus Source List (language only)', 'url': SCOPUS_URL, 'row': row_number, 'status': row[4], 'languages': langs})
        for row_number, row in active:
            indexes.add('Scopus')
            langs = [language_names.get(v.strip(), v.strip()) for v in str(row[7] or '').split(';') if v.strip()]
            languages.update(langs)
            proof = {'index': 'Scopus', 'source': 'Elsevier Scopus Source List', 'snapshot_date': '2026-08', 'retrieved_at': '2026-10-02', 'url': SCOPUS_URL, 'source_id': str(row[0]), 'title': row[1], 'issns': sorted({issn(row[2]), issn(row[3])} - {''}), 'sheet': scopus.worksheets[0].title, 'row': row_number, 'status': row[4], 'coverage': row[5]}
            proofs.append(proof)
            if langs:
                language_proofs.append({'source': proof['source'], 'url': SCOPUS_URL, 'row': row_number, 'languages': langs})
        wos_matches = {n: row for identifier in identifiers for n, row in wos_by_issn.get(identifier, [])}
        for row_number, row in wos_matches.items():
            edition = str(row[4] or '').strip()
            if edition not in {'SSCI', 'ESCI'}:
                continue
            indexes.add(edition)
            proofs.append({'index': edition, 'source': 'UEFISCDI published JCR 2025 list (June 2026 release)', 'snapshot_date': '2026-06', 'retrieved_at': '2026-10-02', 'url': WOS_URL, 'title': row[0], 'issns': sorted({issn(row[1]), issn(row[2])} - {''}), 'category': row[3], 'sheet': wos.worksheets[0].title, 'row': row_number})
        if indexes:
            entries[journal['id']] = {'languages': sorted(languages), 'indexes': [v for v in ['SSCI', 'ESCI', 'Scopus'] if v in indexes], 'evidence': proofs, 'language_evidence': language_proofs}
    return {'schema_version': 1, 'generated_at': '2026-10-02', 'matching': 'Exact current ISSN/eISSN only; no title inference or historical-ISSN inheritance. Scopus requires Active Journal and no discontinued-title record. Language evidence is unioned with existing catalog; any Chinese declaration excludes.', 'sources': [{'url': SCOPUS_URL, 'sha256': hashlib.sha256(scopus_path.read_bytes()).hexdigest(), 'snapshot_date': '2026-08'}, {'url': WOS_URL, 'sha256_readable': hashlib.sha256(wos_path.read_bytes()).hexdigest(), 'snapshot_date': '2026-06'}], 'ambiguous_scopus_ids_excluded': ambiguous, 'excluded_journal_ids': {identifier: 'identity_ambiguity' for identifier in ambiguous}, 'journals': entries}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--scopus', required=True, type=Path)
    parser.add_argument('--wos', required=True, type=Path)
    parser.add_argument('--journals', type=Path, help='Optional full candidate catalog; by default merge the preserved source snapshot and catalog')
    parser.add_argument('--output', type=Path, default=Path('data/catalog/index-evidence.json'))
    parser.add_argument('--language-overrides', type=Path, default=Path('data/catalog/sources/publisher-language-evidence.json'))
    args = parser.parse_args()
    if args.journals:
        candidates = json.loads(args.journals.read_text())
    else:
        from radar_pipeline import merge_journal_catalog
        candidates = merge_journal_catalog(json.loads(Path('data/radar/source_workbook_snapshot.json').read_text()), json.loads(Path('data/catalog/education-journals.json').read_text()))
    payload = build(candidates, args.scopus, args.wos)
    if args.language_overrides.exists():
        overrides = json.loads(args.language_overrides.read_text())
        for journal_id, proof in overrides.items():
            if journal_id not in payload['journals'] or journal_id in payload['excluded_journal_ids']:
                continue
            if not proof.get('url', '').startswith('https://') or not proof.get('excerpt'):
                raise ValueError('Language override requires an explicit publisher statement')
            entry = payload['journals'][journal_id]
            entry['languages'] = sorted(set(entry['languages'] + proof['languages']))
            entry['language_evidence'].append(proof)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'matched_journals': len(payload['journals']), 'index_counts_before_language_filter': dict(Counter(i for j in payload['journals'].values() for i in j['indexes'])), 'ambiguous_scopus_matches': payload['ambiguous_scopus_ids_excluded']}, ensure_ascii=False))
