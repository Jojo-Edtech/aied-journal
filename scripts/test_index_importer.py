"""Offline CLI regressions for full-catalog rebuilding and ambiguous identity."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest


IMPORTER = Path(__file__).with_name("build-index-evidence.py").resolve()


@unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional openpyxl source-reader dependency is unavailable")
class IndexImporterTests(unittest.TestCase):
    def build_fixture(self, root, *, ambiguous=False):
        import openpyxl

        (root / "data/radar").mkdir(parents=True)
        (root / "data/catalog").mkdir(parents=True)
        kept = {"id": "journal-kept", "name": "Existing Journal", "issn": "1234-5678", "languages": ["English"]}
        candidate = {"id": "journal-candidate", "name": "Previously Unverified Journal", "issn": "2345-6789", "languages": []}
        if ambiguous:
            candidate["issns"] = ["2345-6789", "3456-7890"]
        (root / "data/radar/source_workbook_snapshot.json").write_text(json.dumps([kept]))
        (root / "data/catalog/education-journals.json").write_text(json.dumps([candidate]))
        # This is intentionally narrower than the preserved candidate sources.
        (root / "data/radar/journals.json").write_text(json.dumps([kept]))
        scopus = openpyxl.Workbook()
        sources = scopus.active
        sources.title = "Scopus Sources Aug. 2026"
        sources.append(["ID", "Title", "ISSN", "eISSN", "Status", "Coverage", "Discontinued", "Language", "Medline", "OA", "Press", "Added", "Type"])
        for source_id, title, issn in [(1, "Existing Journal", "12345678"), (2, "Previously Unverified Journal", "23456789")]:
            sources.append([source_id, title, issn, "", "Active", "2020-2026", "", "ENG", "", "", "", "", "Journal"])
        if ambiguous:
            sources.append([3, "Different Journal", "34567890", "", "Active", "2020-2026", "", "ENG", "", "", "", "", "Journal"])
        scopus.create_sheet("Accepted Titles")
        discontinued = scopus.create_sheet("Discontinued Titles")
        discontinued.append(["Status: August 2026"])
        discontinued.append(["ID", "Title", "ISSN"])
        scopus.save(root / "scopus.xlsx")
        wos = openpyxl.Workbook()
        wos.active.title = "AIS2025"
        wos.active.append(["Journal name", "ISSN", "eISSN", "Category", "Edition"])
        wos.active.append(["Previously Unverified Journal", "23456789", "", "Education", "ESCI"])
        wos.save(root / "wos.xlsx")

    def run_importer(self, root, *extra):
        subprocess.run(
            [sys.executable, str(IMPORTER), "--scopus", "scopus.xlsx", "--wos", "wos.xlsx", *extra],
            cwd=root, check=True, capture_output=True, text=True,
        )
        return json.loads((root / "data/catalog/index-evidence.json").read_text())

    def test_default_rebuild_reconsiders_candidates_absent_from_public_scope(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.build_fixture(root)
            result = self.run_importer(root)
            self.assertEqual(set(result["journals"]), {"journal-kept", "journal-candidate"})
            candidate = result["journals"]["journal-candidate"]
            self.assertEqual(candidate["languages"], ["English"])
            self.assertEqual(candidate["indexes"], ["ESCI", "Scopus"])
            explicit = self.run_importer(root, "--journals", "data/radar/journals.json")
            self.assertEqual(set(explicit["journals"]), {"journal-kept"})

    def test_cross_journal_issn_ambiguity_is_excluded_before_wos_labels(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.build_fixture(root, ambiguous=True)
            result = self.run_importer(root)
            self.assertNotIn("journal-candidate", result["journals"])
            self.assertEqual(result["excluded_journal_ids"], {"journal-candidate": "identity_ambiguity"})


if __name__ == "__main__":
    unittest.main()
